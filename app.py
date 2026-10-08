from flask import Flask, render_template, request, redirect, session
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

app = Flask(__name__)
app.secret_key = "employee-work-secret"

EXCEL_FILE = "work_data.xlsx"


# =========================
# EXCEL SETUP
# =========================

def get_excel_sheet():

    if not os.path.exists(EXCEL_FILE):

        wb = Workbook()
        ws = wb.active
        ws.title = "Daily Work"

        ws.append([
            "Date",
            "Employee Name",
            "Department",
            "Today's Work",
            "Calls Made",
            "Leads Generated",
            "Follow-ups",
            "Working Hours",
            "Status",
            "Remarks"
        ])

        wb.save(EXCEL_FILE)

        return wb, ws

    wb = load_workbook(EXCEL_FILE)

    # Daily Work sheet असेल तर ती वापरा
    if "Daily Work" in wb.sheetnames:

        ws = wb["Daily Work"]

    else:

        # Existing पहिली sheet वापरा
        ws = wb[wb.sheetnames[0]]

        # तिचे नाव Daily Work करा
        ws.title = "Daily Work"

        # जर sheet रिकामी असेल तर headings तयार करा
        if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:

            ws.append([
                "Date",
                "Employee Name",
                "Department",
                "Today's Work",
                "Calls Made",
                "Leads Generated",
                "Follow-ups",
                "Working Hours",
                "Status",
                "Remarks"
            ])

        wb.save(EXCEL_FILE)

    return wb, ws


# Excel तयार/तपासा
get_excel_sheet()


# =========================
# EMPLOYEE FORM
# =========================

@app.route("/")
def home():
    return render_template("index.html")


# =========================
# SUBMIT WORK
# =========================

@app.route("/submit", methods=["POST"])
def submit():

    wb, ws = get_excel_sheet()

    ws.append([
        datetime.now().strftime("%Y-%m-%d"),
        request.form.get("employee_name", ""),
        request.form.get("department", ""),
        request.form.get("work", ""),
        request.form.get("calls", ""),
        request.form.get("leads", ""),
        request.form.get("followups", ""),
        request.form.get("hours", ""),
        request.form.get("status", ""),
        request.form.get("remarks", "")
    ])

    wb.save(EXCEL_FILE)

    return redirect("/")


# =========================
# MANAGER LOGIN
# =========================

@app.route("/manager-login", methods=["GET", "POST"])
def manager_login():

    if request.method == "POST":

        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "").strip()

        if username == "manager1" and password == "1234":

            session["manager_logged_in"] = True
            session["manager_name"] = "Manager 1"

            return redirect("/dashboard")

        elif username == "manager2" and password == "5678":

            session["manager_logged_in"] = True
            session["manager_name"] = "Manager 2"

            return redirect("/dashboard")

        else:

            return render_template(
                "manager_login.html",
                error="Invalid username or password"
            )

    return render_template("manager_login.html")


# =========================
# DASHBOARD
# =========================

@app.route("/dashboard")
def dashboard():

    if not session.get("manager_logged_in"):

        return redirect("/manager-login")

    wb, ws = get_excel_sheet()

    work_data = list(
        ws.iter_rows(
            min_row=2,
            values_only=True
        )
    )

    total_entries = len(work_data)

    marketing_count = 0
    calling_count = 0

    for row in work_data:

        if len(row) > 2:

            department = str(row[2]).strip().lower()

            if department == "marketing":
                marketing_count += 1

            elif department == "calling":
                calling_count += 1

    return render_template(
        "dashboard.html",
        work_data=work_data,
        total_entries=total_entries,
        marketing_count=marketing_count,
        calling_count=calling_count,
        manager_name=session.get("manager_name")
    )


# =========================
# LOGOUT
# =========================

@app.route("/manager-logout")
def manager_logout():

    session.clear()

    return redirect("/manager-login")


# =========================
# RUN
# =========================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )